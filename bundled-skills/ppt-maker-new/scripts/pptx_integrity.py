#!/usr/bin/env python3
"""Inspect exported OOXML without an Office renderer; never claim visual acceptance."""
import argparse
import hashlib
import zipfile
from pathlib import Path
from xml.etree import ElementTree
from pptx import Presentation
from audience_copy import xml_copy_findings, copy_findings
from stage_runtime import atomic_write_json
from quality_policy import partition


def inspect_pptx(path: Path, expected_slides: int, run_id: str) -> dict:
    errors = []
    count = 0
    digest = None
    try:
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        with zipfile.ZipFile(path) as archive:
            bad = archive.testzip()
            if bad:
                raise ValueError('corrupt ZIP member: ' + bad)
            for name in archive.namelist():
                if name.endswith(('.xml', '.rels')):
                    ElementTree.fromstring(archive.read(name))
        presentation = Presentation(path)
        count = len(presentation.slides)
        if count != expected_slides or expected_slides < 1:
            errors.append({'code': 'slide_count_mismatch', 'message': f'{count} != {expected_slides}'})
        # Read actual editable slide text, including split runs, groups and tables.
        # Speaker notes are deliberately outside audience-visible copy.
        for number, slide in enumerate(presentation.slides, 1):
            errors.extend(xml_copy_findings(slide.element, number, include_advisory=True))
            for shape in slide.shapes:
                if shape.has_chart:
                    chart = ElementTree.fromstring(shape.chart.part.blob)
                    errors.extend(xml_copy_findings(chart, number, include_advisory=True))
                    ns = {'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart'}
                    for node in chart.findall('.//c:strCache/c:pt/c:v', ns):
                        errors.extend(copy_findings(node.text or '', number, 'chart_label', include_advisory=True))
        # Force relationship resolution, including media, instead of only opening the ZIP.
        for part in presentation.part.package.iter_parts():
            for relationship in part.rels.values():
                if not relationship.is_external:
                    _ = relationship.target_part.blob
    except Exception as exc:
        errors.append({'code': 'invalid_pptx', 'message': str(exc)})
    errors, warnings = partition(errors)
    return {'version': 1, 'run_id': run_id, 'input_sha256': digest, 'warnings': warnings,
            'expected_slides': expected_slides, 'slide_count': count,
            'status': 'failed' if errors else 'passed', 'errors': errors,
            'visual_review_status': 'not_performed'}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input', required=True)
    parser.add_argument('--expected-slides', type=int, required=True)
    parser.add_argument('--run-id', required=True)
    parser.add_argument('--report', required=True)
    args = parser.parse_args()
    report = inspect_pptx(Path(args.input), args.expected_slides, args.run_id)
    atomic_write_json(Path(args.report), report)
    return 0 if report['status'] == 'passed' else 1


if __name__ == '__main__':
    raise SystemExit(main())
