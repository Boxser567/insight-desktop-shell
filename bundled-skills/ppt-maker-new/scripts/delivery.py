"""Hash-bound working-deck delivery, separate from strict quality certification.

No local business-content/layout repair: exact SOL pages or a fixed technical
placeholder. Failed drafts never enter validated authoring or success caches.
"""
from __future__ import annotations
import json
import os
import re
import tempfile
from pathlib import Path
from xml.etree import ElementTree as ET
from xml.sax.saxutils import escape

from stage_runtime import atomic_write_json, atomic_write_text, sha256_file, sha256_text, skill_sha256, render_contract_sha256
from pptx_integrity import inspect_pptx
from quality_policy import blocking, partition, height_estimate_risk


def read(path):
    try:
        value = json.loads(path.read_text(encoding='utf-8'))
        return value if isinstance(value, dict) else {}
    except (OSError, ValueError):
        return {}


def usable_imperfect(xml: str, errors: list[dict]) -> bool:
    """Allow bounded estimated fit risk, never content/schema/occlusion failures."""
    if any(height_estimate_risk(e) != 2 for e in blocking(errors)):
        return False
    try:
        root = ET.fromstring(xml)
        # A small typographic shortfall is tolerated, never unreadable tiny copy.
        return all(float(n.get('fontSize', '0')) >= (10 if n.get('textType') == 'caption' else 13)
                   for n in root.iter() if n.tag.split('}')[-1] == 'content')
    except (ValueError, ET.ParseError):
        return False


def placeholder_xml(page: dict) -> str:
    from audience_copy import copy_findings
    title = str(page.get('title') or f"第 {page['no']} 页")
    if copy_findings(title, int(page['no']), 'title'):
        title = f"第 {page['no']} 页"
    # Fixed status artifact explicitly authorized by the delivery policy, not
    # a template compiler for normal pages. Never invent missing business copy.
    return f'''<slide xmlns="https://www.larkoffice.com/sml/2.0"><data>
<shape type="text" topLeftX="60" topLeftY="70" width="840" height="210"><content textType="title" fontSize="28" wrap="true"><p>{escape(title)}</p></content></shape>
<shape type="text" topLeftX="60" topLeftY="330" width="840" height="80"><content textType="body" fontSize="20" wrap="true"><p>本页暂未完成，保留原页位置。其余页面可先使用。</p></content></shape>
<shape type="text" topLeftX="60" topLeftY="470" width="840" height="30"><content textType="caption" fontSize="12"><p>第 {int(page['no'])} 页 · 待完善版本 · 未进行视觉验收</p></content></shape>
</data></slide>'''


def candidate_rank(xml: str, errors: list[dict]) -> tuple[int, int, int]:
    """Prefer the least risky saved model draft; never compute replacement SML."""
    rank = 0 if not errors else 1 if usable_imperfect(xml, errors) else 2
    advisory = 0
    if rank == 0:
        from xml_lint import extract_elements, detect_text_may_overflow_shapes, detect_text_may_wrap_shapes
        elements = extract_elements(xml)
        advisory = sum(issue.get('level') == 'info' for issue in
                       detect_text_may_overflow_shapes(elements) + detect_text_may_wrap_shapes(elements))
    return rank, len(errors), advisory


def _presentation():
    from pptx import Presentation
    from sml_to_pptx import px
    prs = Presentation()
    prs.slide_width, prs.slide_height = px(960), px(540)
    return prs


def _render(prs, xml, assets):
    from sml_to_pptx import render_slide
    try:
        from iconpark_render import icon_renderer
    except ImportError:
        icon_renderer = None
    warnings = []
    render_slide(prs, ET.fromstring(xml), {}, assets, None, icon_renderer, warnings)
    errors, advisory = partition(warnings)
    if errors:
        raise ValueError('converter blocking findings: ' + json.dumps(errors, ensure_ascii=False))
    return advisory


def _retain_external_findings(findings, authoring):
    """Replace only verifiable lint-owned evidence with fresh candidate lint.

    A page number or error code alone is not authority to waive an external
    finding. Bind to the actual source bytes and confirm that the same validator
    reports that issue there. Unknown, untagged and dependency evidence remains.
    """
    if not authoring.is_file() or not any(e.get('validation_source') == 'sml_lint' for e in findings):
        return findings
    from sol_render import lint_slide_xml
    old_xml = authoring.read_text(encoding='utf-8')
    digest = sha256_text(old_xml)
    warnings = []
    old_findings = lint_slide_xml(old_xml, warnings_out=warnings) + warnings
    def identity(e):
        return (e.get('code'), tuple(e.get('elements') or ()))
    reproduced = {identity(e) for e in old_findings}
    return [e for e in findings if not (e.get('validation_source') == 'sml_lint'
            and e.get('sml_sha256') == digest and identity(e) in reproduced)]


def export_best_effort(project: Path, output: Path, *, run_id: str,
                       extra_findings: list[dict] | None = None,
                       candidate_overrides: dict[int, str] | None = None,
                       promote_candidates: bool = True, recover_pending: bool = True) -> dict:
    from sol_render import _prepare_render_state, _validate_batch_candidate, _write_render_success
    from render_checkpoint import load_candidates
    project, output = Path(project).resolve(), Path(output).resolve()
    if recover_pending:
        from final_repair import recover_pending_transaction
        recover_pending_transaction(project, output)
    if not re.fullmatch(r'[A-Za-z0-9._-]{1,96}', run_id) or run_id in {'.','..'}:
        raise ValueError('invalid run identity')
    run = project / 'reports/runs' / run_id
    contract = read(run/'workflow_contract.json')
    if contract.get('delivery_policy') != 'best-effort':
        raise RuntimeError('strict/legacy workflow does not permit degraded delivery')
    if (project/'reports/cancel'/f'{run_id}.json').exists():
        raise RuntimeError('cancelled run cannot auto-deliver')
    bp_path = project/'reports/blueprint.json'
    from plan_request import verify_plan_request
    verify_plan_request(project)
    blueprint = read(bp_path)
    pages = blueprint.get('pages') or []
    if [p.get('no') for p in pages] != list(range(1,len(pages)+1)) or not pages:
        raise RuntimeError('delivery requires a complete ordered blueprint')
    if (project/'reports/handoff_receipt.json').exists():
        from import_handoff import verify_receipt
        verify_receipt(project)
    original_blueprint_hash = sha256_file(bp_path)
    from sol_build import _authoring_sha256
    skill_root=Path(__file__).resolve().parents[1]
    current_skill_hash=skill_sha256(skill_root)
    current_render_hash=render_contract_sha256(skill_root)
    from sol_common import resolve_asset_target
    rows, chosen = [], []
    extra_findings = list(extra_findings or [])
    # Offline delivery must not bless changed tracked assets by simply hashing
    # their new bytes. Recheck existing lineage before selecting any page.
    preflight_path=project/'reports/asset_preflight.json'
    preflight_hash=sha256_file(preflight_path) if preflight_path.exists() else None
    if preflight_path.exists():
        from asset_preflight import verify_assets
        try:
            preflight=json.loads(preflight_path.read_text())
            asset_issues=list(preflight.get('errors') or [])
            for row in preflight.get('assets',[]):
                failures=verify_assets([row])
                consumers=[p for p in pages if any(i.get('src')==row.get('src') for i in p.get('images',[]))]
                if consumers:
                    try:
                        if sha256_file(resolve_asset_target(project/'assets',row['src'])) != row['output_sha256']:
                            failures.append('actual consumer asset changed')
                    except (OSError,ValueError,KeyError):
                        failures.append('actual consumer asset unavailable')
                    if failures:
                        asset_issues.append({'code':'asset_consumer_error','src':row['src'],'message':'; '.join(failures)})
            for error in asset_issues:
                consumers=[p['no'] for p in pages if error.get('src') and any(
                    i.get('src')==error['src'] for i in p.get('images',[]))]
                if consumers:
                    extra_findings.extend(dict(error,page_no=n,page=n) for n in consumers)
                elif not error.get('src'):
                    raise RuntimeError('global asset preflight failure: '+str(error))
        except (OSError,ValueError,KeyError,TypeError,AttributeError) as exc:
            raise RuntimeError('invalid asset preflight evidence: '+str(exc)) from exc
    global_findings = [e for e in extra_findings if not (e.get('page') or e.get('page_no'))]
    # Deck-wide quality limitations remain disclosed; authority/provenance failures
    # are never converted into a cosmetic warning or hidden by a placeholder.
    if any(any(word in str(e.get('code','')) for word in
               ('provenance','receipt','hash','authority','must_keep','blueprint','source')) for e in global_findings):
        raise RuntimeError('global content/provenance failure cannot be downgraded')
    for page in pages:
        style = ''
        if not (blueprint.get('theme') or {}).get('render_contract'):
            from sol_plan import select_style_reference
            cache=read(project/'reports/render_attempts'/f"slide-{page['no']:02d}-cache.json")
            if 'legacy_style_text' in cache:
                style=cache['legacy_style_text']
            else:
                style=select_style_reference(read(project/'reports/brief.json'),Path(__file__).resolve().parents[1]).read_text()
                for path in (project/'reports/render_checkpoints').glob('*.json'):
                    old=read(path)
                    if old.get('page_no')==page['no'] and 'legacy_style_text' in old:
                        style=old['legacy_style_text']; break
        errors = [e for e in extra_findings if (e.get('page') or e.get('page_no')) == page['no']]
        # Bind before state preparation can restore an immutable accepted artifact.
        errors = _retain_external_findings(errors, project/'authoring'/f"slide-{page['no']:02d}.xml")
        blocked_asset = any(e.get('code') in {'asset_consumer_error','asset_dependency_failed'} for e in errors)
        state = _prepare_render_state(project,blueprint,page,style,not blocked_asset,current_skill_hash,current_render_hash)
        cached_xml = state['output'].read_text(encoding='utf-8') if state['cached'] else None
        options = []
        if not blocked_asset:
            override = (candidate_overrides or {}).get(page['no'])
            if override:
                options.append((override, 'final_repair_candidate'))
            if cached_xml:
                options.append((cached_xml, 'validated_authoring'))
            options.extend((draft, 'model_checkpoint') for draft in load_candidates(state))
        ranked, seen = [], set()
        for draft, origin in options:
            if draft in seen:
                continue
            seen.add(draft)
            candidate, findings, _ = _validate_batch_candidate(state, draft, check_conversion=False)
            combined, advisory = partition(errors + findings)
            candidate_warnings = list(state.get('quality_warnings') or []) + advisory
            if not combined or usable_imperfect(candidate, combined):
                # Imperfect candidates also need backend preflight. Try the next
                # exact candidate if this one cannot be converted; no model call.
                from page_conversion import conversion_findings
                converted_errors, converted_warnings = partition(conversion_findings(candidate, project/'assets', renderer=_render))
                combined += converted_errors
                candidate_warnings += converted_warnings
            ranked.append((candidate_rank(candidate, combined), candidate, combined, candidate_warnings, origin))
        if ranked:
            _, xml, errors, quality_warnings, source = min(ranked, key=lambda row: row[0])
        else:
            xml, source, quality_warnings = None, 'technical_placeholder', []
            errors.append({'code':'missing_slide','message':'No current usable model-authored page'})
        errors, advisory = partition(errors)
        quality_warnings += advisory
        status = 'validated' if xml and not errors else 'imperfect'
        if not xml or errors and not usable_imperfect(xml, errors):
            status = 'placeholder'
            xml = placeholder_xml(page)
        if status == 'validated' and xml != cached_xml and promote_candidates:
            state['quality_warnings'] = quality_warnings
            _write_render_success(state,xml)
        row = {'page_no':page['no'], 'status':status, 'source':source if status!='placeholder' else 'technical_placeholder',
               'prompt_sha256':state['prompt_hash'], 'sml_sha256':sha256_text(xml), 'errors':errors, 'quality_warnings':quality_warnings,
               'assets':({i['src']:sha256_file(resolve_asset_target(project/'assets',i['src']))
                          for i in page.get('images',[])} if status!='placeholder' else {})}
        rows.append(row); chosen.append(xml)
    selection = run/'delivery_slides'
    selection.mkdir(parents=True, exist_ok=True)
    prs = _presentation()
    prs.core_properties.title = str(blueprint.get('topic') or 'PPT') + ' · 待完善版'
    for row,xml in zip(rows,chosen):
        path = selection/f"slide-{row['page_no']:02d}.xml"
        atomic_write_text(path,xml)
        row['selected_path'] = str(path.relative_to(project))
        _render(prs,xml,project/'assets')
        if row['status'] != 'validated':
            prs.slides[-1].notes_slide.notes_text_frame.text = '待完善页面；详见交付问题清单。未进行视觉验收。'
    output.parent.mkdir(parents=True,exist_ok=True)
    fd, temporary = tempfile.mkstemp(prefix='.working-deck-',suffix='.pptx',dir=output.parent)
    os.close(fd)
    temp = Path(temporary)
    try:
        prs.save(temp)
        integrity = inspect_pptx(temp,len(pages),run_id)
        if integrity['status'] != 'passed':
            raise RuntimeError('working deck failed integrity: '+str(integrity['errors']))
        if sha256_file(bp_path) != original_blueprint_hash:
            raise RuntimeError('blueprint changed while selecting delivery pages')
        for row in rows:
            for src,digest in row['assets'].items():
                if sha256_file(resolve_asset_target(project/'assets',src)) != digest:
                    raise RuntimeError('selected asset changed during assembly')
        if skill_sha256(skill_root) != current_skill_hash:
            raise RuntimeError('skill changed during assembly')
        if (sha256_file(preflight_path) if preflight_path.exists() else None) != preflight_hash:
            raise RuntimeError('asset preflight evidence changed during assembly')
        atomic_write_json(run/'delivery_integrity.json',integrity)
        # One manifest commit point follows complete PPTX generation. A crash
        # between replacement and manifest publication produces a hash mismatch,
        # never a falsely approved deliverable.
        os.replace(temp,output)
    finally:
        temp.unlink(missing_ok=True)
    report = {'version':1,'run_id':run_id,'status':'needs_attention','delivery_policy':'best-effort',
              'strict_quality_passed':False,'visual_review_status':'not_performed',
              'page_count':len(pages),'pages':rows,'global_findings':global_findings,
              'attention_pages':[r['page_no'] for r in rows if r['status']!='validated'],
              'validated_pages':sum(r['status']=='validated' for r in rows),
              'output_path':str(output),'output_sha256':sha256_file(output),
              'blueprint_sha256':original_blueprint_hash,'authoring_sha256':_authoring_sha256(project),
              'workflow_sha256':sha256_file(run/'workflow_contract.json'),
              'integrity_sha256':sha256_file(run/'delivery_integrity.json'),
              'asset_preflight_sha256':preflight_hash,
              'skill_sha256':current_skill_hash}
    from delivery_quality import quality_summary, quality_sentence
    from qa_gate import evaluate
    # This audits the actual delivered selection, not stale failed authoring or
    # previous-run QA. Local only; it does not fabricate strict render provenance.
    selected_qa = evaluate(project, read(project/'reports/brief.json'), blueprint, selected_dir=selection)
    report['quality'] = quality_summary(rows, global_findings, selected_qa)
    report['quality_message'] = quality_sentence(report['quality'])
    atomic_write_json(run/'delivery_manifest.json',report)
    atomic_write_json(project/'reports/delivery_manifest.json',report)
    note = ['# 待完善交付清单','',f"共 {len(pages)} 页；{report['validated_pages']} 页通过单页检查。",
            report['quality_message'],'']
    note += [f"- 第 {r['page_no']} 页：{r['status']}；"+', '.join(str(e.get('code')) for e in r['errors']) for r in rows if r['status']!='validated']
    note += [f"- 第 {r['page_no']} 页非阻断提醒："+', '.join(str(w.get('code')) for w in r['quality_warnings']) for r in rows if r['quality_warnings']]
    if global_findings: note.append('- 全稿限制：'+json.dumps(global_findings,ensure_ascii=False))
    if selected_qa['errors']: note.append('- 当前整稿检查：'+json.dumps(selected_qa['errors'], ensure_ascii=False))
    atomic_write_text(project/'reports/delivery_issues.md','\n'.join(note)+'\n')
    from sol_build import write_build_state
    write_build_state(project,'needs_attention',run_id=run_id,skill_sha256=report['skill_sha256'])
    # Never issue success/release_ready for this separate evidence chain.
    marker = {'run_id':run_id,'status':'needs_attention','delivery_manifest_sha256':sha256_file(run/'delivery_manifest.json')}
    atomic_write_json(run/'release_manifest.json',marker)
    atomic_write_json(project/'reports/release_manifest.json',marker)
    return report


def verify_working_delivery(project: Path, output: Path) -> list[str]:
    from sol_build import _authoring_sha256
    errors = []
    marker = read(project/'reports/release_manifest.json')
    run_id = marker.get('run_id','')
    if not re.fullmatch(r'[A-Za-z0-9._-]{1,96}',run_id) or run_id in {'.','..'}:
        return ['invalid delivery run identity']
    run = project/'reports/runs'/run_id
    report = read(run/'delivery_manifest.json')
    try:
        if marker.get('status') != 'needs_attention' or report.get('status') != 'needs_attention':
            errors.append('not a working-deck delivery')
        if marker.get('delivery_manifest_sha256') != sha256_file(run/'delivery_manifest.json'):
            errors.append('delivery manifest changed')
        if read(run/'workflow_contract.json').get('delivery_policy') != 'best-effort':
            errors.append('strict policy forbids working-deck publication')
        if report.get('workflow_sha256') != sha256_file(run/'workflow_contract.json'):
            errors.append('workflow changed')
        if report.get('skill_sha256') != skill_sha256(Path(__file__).resolve().parents[1]):
            errors.append('skill changed')
        if report.get('output_sha256') != sha256_file(output): errors.append('PPTX changed')
        if report.get('blueprint_sha256') != sha256_file(project/'reports/blueprint.json'): errors.append('blueprint changed')
        if report.get('authoring_sha256') != _authoring_sha256(project): errors.append('authoring changed')
        if report.get('integrity_sha256') != sha256_file(run/'delivery_integrity.json'): errors.append('integrity evidence changed')
        preflight_path=project/'reports/asset_preflight.json'
        if report.get('asset_preflight_sha256') != (sha256_file(preflight_path) if preflight_path.exists() else None):
            errors.append('asset preflight evidence changed')
        pages = report.get('pages') or []
        if [r['page_no'] for r in pages] != list(range(1,report.get('page_count',0)+1)) or not pages:
            errors.append('incomplete page selection')
        for row in pages:
            path = (project/row['selected_path']).resolve()
            if not path.is_relative_to((run/'delivery_slides').resolve()) or sha256_file(path)!=row['sml_sha256']:
                errors.append('selected SML changed')
            from sol_common import resolve_asset_target
            for src,digest in row.get('assets',{}).items():
                if sha256_file(resolve_asset_target(project/'assets',src)) != digest:
                    errors.append('selected asset changed')
        integrity = inspect_pptx(output,report['page_count'],run_id)
        if integrity['status']!='passed': errors.append('PPTX integrity failed')
        if (project/'reports/handoff_receipt.json').exists():
            from import_handoff import verify_receipt
            verify_receipt(project)
        build = read(project/'reports/build_state.json')
        if build.get('run_id') != run_id:
            errors.append('delivery is not from current build')
        if build.get('status') != 'needs_attention' or build.get('skill_sha256') != report.get('skill_sha256'):
            errors.append('working build status or skill identity changed')
        if (project/'reports/cancel'/f'{run_id}.json').exists(): errors.append('cancelled run')
    except (OSError,ValueError,KeyError,TypeError,RuntimeError) as exc:
        errors.append('invalid delivery evidence: '+str(exc))
    return errors


def main():
    import argparse
    parser=argparse.ArgumentParser(description='Offline delivery of current pages; never starts model work')
    parser.add_argument('--project',required=True)
    parser.add_argument('--run-id',required=True)
    parser.add_argument('--delivery-policy',choices=['best-effort','strict'],default='best-effort')
    args=parser.parse_args()
    project=Path(args.project).resolve()
    if args.delivery_policy=='strict': raise RuntimeError('strict mode requires normal successful build')
    from blueprint_validation import static_findings
    findings=static_findings(read(project/'reports/blueprint.json'),read(project/'reports/brief.json'))
    if findings: raise RuntimeError('blueprint is incomplete or invalid: '+str(findings))
    from sol_build import _claim_run_id
    _claim_run_id(project,args.run_id)
    atomic_write_json(project/'reports/runs'/args.run_id/'workflow_contract.json',
                      {'version':8,'delivery_policy':args.delivery_policy,'visual_mode':'off','offline_delivery_only':True})
    report=export_best_effort(project,project/'out.pptx',run_id=args.run_id)
    print(json.dumps(report,ensure_ascii=False))
    return 0


if __name__=='__main__':
    raise SystemExit(main())
