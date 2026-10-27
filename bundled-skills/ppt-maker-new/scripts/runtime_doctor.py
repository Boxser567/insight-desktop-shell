"""Offline dependency and host smoke checks. No model requests or visual rendering."""
import importlib
from importlib.metadata import version, PackageNotFoundError
import os
from pathlib import Path
import subprocess
import sys
import tempfile

from platform_runtime import acquire_lock, child_environment, ADAPTER_VERSION


def diagnose(project):
    from sol_common import transport_config
    from ppt_proxy import proxy_environment_ready
    result = {'python':sys.version.split()[0], 'platform':sys.platform,
              'host_adapter':ADAPTER_VERSION, 'network_checked':False,
              'visual_review_supported':False, 'dependencies':{}, 'dependency_versions':{},
              'local_checks':{}, 'gateway':transport_config(),
              'blueprint_present':(project/'reports/blueprint.json').is_file(),
              'managed_retention':'requires_client_acceptance_test',
              'multimodal_gateway':'not_checked'}
    try:
        result['enterprise_proxy_ready'] = proxy_environment_ready()
    except (OSError, RuntimeError, ValueError):
        result['enterprise_proxy_ready'] = False
    result['server_credentials'] = 'not_checked'
    required = {'pptx':'python-pptx', 'PIL':'Pillow', 'lxml':'lxml', 'fontTools':'fonttools'}
    if sys.platform == 'win32': required['psutil'] = 'psutil'
    for name, distribution in required.items():
        try:
            importlib.import_module(name)
            result['dependencies'][name] = True
            try: result['dependency_versions'][name] = version(distribution)
            except PackageNotFoundError: result['dependency_versions'][name] = 'unknown'
        except Exception as exc:
            result['dependencies'][name] = False
            result['dependency_versions'][name] = type(exc).__name__
    def check(name, operation):
        try:
            operation()
            result['local_checks'][name] = {'ok':True}
        except Exception as exc:
            result['local_checks'][name] = {'ok':False, 'error_type':type(exc).__name__}
    def smoke():
        project.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(prefix='.ppt-doctor-',dir=project) as td:
            root=Path(td)
            def files():
                from stage_runtime import atomic_write_text
                path=root/'中文 #1.txt'
                atomic_write_text(path,'中文 🎯\n')
                destination=root/'已移动.txt'
                os.replace(path,destination)
                if destination.read_text(encoding='utf-8') != '中文 🎯\n':
                    raise ValueError('UTF-8 roundtrip mismatch')
            check('utf8_atomic_files',files)
            def pptx():
                from pptx import Presentation
                from pptx.util import Inches
                prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6])
                slide.shapes.add_textbox(Inches(1),Inches(1),Inches(4),Inches(1)).text='中文 ABC 123'
                path=root/'中文.pptx';prs.save(path)
                if Presentation(path).slides[0].shapes[0].text != '中文 ABC 123':
                    raise ValueError('PPTX text roundtrip mismatch')
            check('pptx_roundtrip',pptx)
            def locking():
                path=root/'owner.lock'
                code=("import sys;from pathlib import Path;from platform_runtime import acquire_lock;"
                      "h=acquire_lock(Path(sys.argv[1]));h.close()")
                environment=child_environment()
                environment['PYTHONPATH']=str(Path(__file__).parent)
                def probe():
                    return subprocess.run([sys.executable,'-c',code,str(path)],env=environment,
                                          capture_output=True,timeout=10)
                handle=acquire_lock(path)
                try:
                    blocked=probe()
                    if blocked.returncode == 0 or b'project already has an active build' not in blocked.stderr:
                        raise RuntimeError('child did not observe the lock')
                finally:handle.close()
                if probe().returncode != 0: raise RuntimeError('released lock remained blocked')
            check('process_lock',locking)
    check('workspace_smoke',smoke)
    try:
        from font_policy import resolve_font_contract, _font_dirs
        fonts, issues=resolve_font_contract({})
        result['fonts']={'resolved':fonts,'issues':issues,
                         'search_directories':[str(p) for p in _font_dirs()],
                         'cross_machine_layout_verified':False}
    except Exception as exc:
        result['fonts']={'issues':[{'code':'font_inspection_failed','error_type':type(exc).__name__}],
                         'cross_machine_layout_verified':False}
    result['local_ready']=(sys.version_info >= (3,10) and all(result['dependencies'].values())
                           and all(row['ok'] for row in result['local_checks'].values()))
    return result
