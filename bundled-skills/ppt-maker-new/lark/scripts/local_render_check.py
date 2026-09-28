#!/usr/bin/env python3
"""Local OOXML structural diagnostics only; no screenshot or visual review."""
from __future__ import annotations
import argparse
import hashlib
import json
import sys
from pathlib import Path
from typing import Any
from pptx import Presentation
from pptx.util import Emu

def structural_check(pptx_path: Path) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    errors: list[dict[str, Any]] = []
    warnings: list[dict[str, Any]] = []
    prs = Presentation(str(pptx_path))
    sw, sh = prs.slide_width, prs.slide_height
    if len(prs.slides) == 0:
        errors.append({"code": "empty_deck", "message": "deck contains no slides"})
        return errors, warnings
    for idx, slide in enumerate(prs.slides, 1):
        shapes = list(slide.shapes)
        if not shapes:
            errors.append({"code": "blank_slide", "message": f"slide {idx} has no shapes"})
            continue
        has_visible = False
        for shape in shapes:
            if shape.left is None:
                continue
            right = shape.left + shape.width
            bottom = shape.top + shape.height
            tolerance = Emu(1000)
            if (
                shape.left < -tolerance
                or shape.top < -tolerance
                or right > sw + tolerance
                or bottom > sh + tolerance
            ):
                errors.append(
                    {
                        "code": "shape_out_of_canvas",
                        "message": f"slide {idx}: {shape.shape_type} exceeds slide bounds "
                        f"(left={shape.left} top={shape.top} right={right} bottom={bottom}; slide {sw}x{sh})",
                    }
                )
            if shape.has_text_frame and shape.text_frame.text.strip():
                has_visible = True
            if shape.shape_type == 13 or shape._element.xpath('.//a:blip'):  # picture or shape picture-fill
                has_visible = True
            if shape.has_table:
                has_visible = True
            if shape.has_chart:
                has_visible = True
        if not has_visible:
            warnings.append(
                {"code": "possibly_blank_slide", "message": f"slide {idx} has shapes but no visible content"}
            )
    return errors, warnings


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input', required=True)
    parser.add_argument('--expected-slides', type=int)
    parser.add_argument('--json', action='store_true')
    parser.add_argument('--report')
    args = parser.parse_args(argv)
    input_path = Path(args.input)
    if not input_path.is_file():
        print(json.dumps({'error': f'input not found: {input_path}'}))
        return 2
    errors, warnings = structural_check(input_path)
    actual = len(Presentation(str(input_path)).slides)
    if args.expected_slides is not None and actual != args.expected_slides:
        errors.append({'code': 'slide_count_mismatch',
                       'message': f'expected {args.expected_slides} slides, got {actual}'})
    report = {
        'tool': 'local_render_check', 'file': str(input_path),
        'input_sha256': hashlib.sha256(input_path.read_bytes()).hexdigest(),
        'visual_review_status': 'not_performed',
        'summary': {'error_count': len(errors), 'warning_count': len(warnings),
                    'structural_check_passed': not errors, 'release_ready': False,
                    'release_stage': 'component_structural_check', 'slide_count': actual},
        'errors': errors, 'warnings': warnings,
    }
    if args.report:
        report_path = Path(args.report)
        report_path.parent.mkdir(parents=True, exist_ok=True)
        report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 1 if errors else 0


if __name__ == '__main__':
    sys.exit(main())
