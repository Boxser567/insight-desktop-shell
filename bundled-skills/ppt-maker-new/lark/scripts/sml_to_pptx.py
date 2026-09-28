#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Convert Feishu Slides SML XML (the lark-slides-pro authoring format) into a
local editable .pptx deck.

Local-first mode of the lark-slides-pro skill: the agent still authors each
page as SML <slide> XML and gates it with xml_lint.py; this converter is the
"write to slides" step replaced by a local render. The Feishu-only constructs
that cannot be represented in pptx are degraded explicitly (see --json
report .warnings) instead of silently dropped.

Usage:
  python3 sml_to_pptx.py --input <deck.xml | dir> --output out.pptx \
      [--assets DIR] [--media-manifest manifest.json] [--title T] [--json]

Input is either:
  * one <presentation> XML file (slides inside), or
  * a directory of slide-XX.xml files (sorted by name) with an optional
    presentation.xml / theme.xml carrying <theme> defaults.

Coordinate model: the SML canvas is 960x540 px at 72 px/inch, so geometry
maps 1 px -> 12700 EMU and font sizes map 1 px -> 1 pt. This matches the
pptx 16:9 slide (13.333 x 7.5 in).
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path
from typing import Any, Optional
from xml.etree import ElementTree as ET

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt
from lxml import etree  # python-pptx's XML tree
sys.path.insert(0, str(Path(__file__).resolve().parents[2] / 'scripts'))
from quality_policy import partition
from backend_capabilities import (SHAPE_ENUM_NAMES, CHART_ENUM_NAMES,
                                  POLYLINE_ENUM_NAMES, POLYLINE_PRESETS, AUTO_FIT_ENUM_NAMES,
                                  capability_findings)

SML_NS = "https://www.larkoffice.com/sml/2.0"
CANVAS_W, CANVAS_H = 960, 540
EMU_PER_PX = 12700  # 914400 EMU/in / 72 px/in

TEXT_TYPE_DEFAULT_SIZE = {
    "title": 54,
    "headline": 38,
    "sub-headline": 32,
    "body": 16,
    "caption": 12,
}

DEFAULT_FONT = "思源黑体"
DEFAULT_TEXT_COLOR = "rgb(31, 35, 41)"

SHAPE_TYPE_MAP = {key: getattr(MSO_SHAPE, name) for key, name in SHAPE_ENUM_NAMES.items()}

LINE_DASH_MAP = {
    "solid": None,
    "dash": "dash",
    "dot": "dot",
    "long-dash": "dash",
    "round-dot": "dot",
}

ARROW_TYPE_MAP = {
    "arrow": "triangle",
    "solid-triangle": "triangle",
    "empty-triangle": "triangle",
    "solid-diamond": "diamond",
    "empty-diamond": "diamond",
    "solid-circle": "oval",
    "empty-circle": "oval",
}

CHART_TYPE_MAP = {key: getattr(XL_CHART_TYPE, name) for key, name in CHART_ENUM_NAMES.items()}

LEGEND_POSITION_MAP = {
    "top": XL_LEGEND_POSITION.TOP,
    "bottom": XL_LEGEND_POSITION.BOTTOM,
    "left": XL_LEGEND_POSITION.LEFT,
    "right": XL_LEGEND_POSITION.RIGHT,
}


# --------------------------------------------------------------------------
# small helpers
# --------------------------------------------------------------------------

def local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def children(el: ET.Element, name: str) -> list[ET.Element]:
    return [c for c in el if local(c.tag) == name]


def dim_children(el: ET.Element) -> list[ET.Element]:
    """chartData dimensions are literally <dim1>, <dim2>, ... in SML."""
    return [c for c in el if re.fullmatch(r"dim\d+", local(c.tag))]


def child(el: ET.Element, name: str) -> Optional[ET.Element]:
    for c in el:
        if local(c.tag) == name:
            return c
    return None


def attr(el: ET.Element, name: str, default: Any = None) -> Any:
    return el.get(name, default)


def px(value: Any) -> Emu:
    return Emu(int(float(value) * EMU_PER_PX))


def pt(value: Any) -> Pt:
    return Pt(float(value))


class ColorSpec:
    """Parsed SML color: solid rgb + alpha, or a two-stop linear gradient."""

    __slots__ = ("rgb", "alpha", "gradient", "kind", "angle")

    def __init__(self, rgb: Optional[RGBColor], alpha: float, gradient: Optional[list[tuple[RGBColor, float]]], kind: str, angle: Optional[float] = None):
        self.rgb = rgb
        self.alpha = alpha
        self.gradient = gradient  # [(rgb, position 0..1), ...]
        self.kind = kind  # solid | linear-gradient | radial-gradient | transparent
        self.angle = angle  # linear-gradient angle in degrees


_RGB_RE = re.compile(r"rgb\(\s*(\d+)\s*,\s*(\d+)\s*,\s*(\d+)\s*\)")
_RGBA_RE = re.compile(r"rgba\(\s*(\d+)\s*,\s*(\d+)\s*,\s*(\d+)\s*,\s*([0-9.]+)\s*\)")


def parse_color(value: Optional[str]) -> Optional[ColorSpec]:
    if not value:
        return None
    value = value.strip()
    if value == "transparent":
        return ColorSpec(None, 0.0, None, "transparent")
    m = _RGBA_RE.match(value)
    if m:
        return ColorSpec(
            RGBColor(int(m.group(1)), int(m.group(2)), int(m.group(3))),
            float(m.group(4)),
            None,
            "solid",
        )
    m = _RGB_RE.match(value)
    if m:
        return ColorSpec(
            RGBColor(int(m.group(1)), int(m.group(2)), int(m.group(3))),
            1.0,
            None,
            "solid",
        )
    # linear-gradient(135deg,rgba(...) 0%,rgba(...) 100%) — the inner contains
    # nested parens, so capture greedily to the LAST ")" then re-parse stops.
    m = re.match(r"(linear|radial)-gradient\((.*)\)\s*$", value, re.IGNORECASE)
    if m:
        kind = f"{m.group(1).lower()}-gradient"  # consumers check "linear-gradient"/"radial-gradient"
        inner = m.group(2)
        angle = None
        angle_m = re.match(r"^\s*([\d.]+)deg\s*,?\s*", inner, re.IGNORECASE)
        if angle_m:
            angle = float(angle_m.group(1))
            inner = inner[angle_m.end():]
        parsed: list[tuple[RGBColor, float]] = []
        for sm in re.finditer(r"(rgba?\([^)]*\))\s*([\d.]+%)?", inner):
            cm = _RGBA_RE.match(sm.group(1)) or _RGB_RE.match(sm.group(1))
            if not cm:
                continue
            rgb = RGBColor(int(cm.group(1)), int(cm.group(2)), int(cm.group(3)))
            pos_str = sm.group(2)
            pos = float(pos_str.rstrip("%")) / 100.0 if pos_str else 0.0
            parsed.append((rgb, pos))
        if len(parsed) == 1:
            parsed.append((parsed[0][0], 1.0))
        if parsed:
            return ColorSpec(None, 1.0, parsed, kind, angle)
    return None


def alpha_of(alpha: Optional[float]) -> float:
    return alpha if alpha is not None else 1.0


# --------------------------------------------------------------------------
# text building
# --------------------------------------------------------------------------

class TextStyle:
    """Effective text style for a run (merged from content + inline span)."""

    __slots__ = (
        "font_size", "font_family", "color", "bold", "italic",
        "underline", "strikethrough", "letter_spacing",
    )

    def __init__(self) -> None:
        self.font_size: Optional[float] = None
        self.font_family: Optional[str] = None
        self.color: Optional[str] = None
        self.bold: Optional[bool] = None
        self.italic: Optional[bool] = None
        self.underline: Optional[bool] = None
        self.strikethrough: Optional[bool] = None
        self.letter_spacing: Optional[float] = None

    def merge(self, other: "TextStyle") -> None:
        for name in self.__slots__:
            value = getattr(other, name)
            if value is not None:
                setattr(self, name, value)


def content_style(content: ET.Element, theme_defaults: dict[str, dict[str, Any]]) -> TextStyle:
    st = TextStyle()
    text_type = attr(content, "textType")
    td = theme_defaults.get(text_type or "body", {}) if text_type else {}
    st.font_size = _num(attr(content, "fontSize"), td.get("fontSize"))
    st.font_family = attr(content, "fontFamily") or td.get("fontFamily") or DEFAULT_FONT
    st.color = attr(content, "color") or td.get("fontColor") or DEFAULT_TEXT_COLOR
    st.bold = _bool(attr(content, "bold"), None)
    st.italic = _bool(attr(content, "italic"), None)
    st.underline = _bool(attr(content, "underline"), None)
    st.strikethrough = _bool(attr(content, "strikethrough"), None)
    st.letter_spacing = _num(attr(content, "letterSpacing"), None)
    if st.font_size is None:
        st.font_size = TEXT_TYPE_DEFAULT_SIZE.get(text_type or "body", 16)
    return st


def _num(value: Any, default: Any = None) -> Any:
    if value is None or value == "":
        return default
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def _bool(value: Any, default: Any = None) -> Any:
    if value is None:
        return default
    return str(value).strip().lower() in ("true", "1")


def span_style(el: ET.Element) -> TextStyle:
    st = TextStyle()
    if attr(el, "fontSize") is not None:
        st.font_size = _num(attr(el, "fontSize"))
    if attr(el, "fontFamily") is not None:
        st.font_family = attr(el, "fontFamily")
    if attr(el, "color") is not None:
        st.color = attr(el, "color")
    st.bold = _bool(attr(el, "bold"), None)
    st.italic = _bool(attr(el, "italic"), None)
    st.underline = _bool(attr(el, "underline"), None)
    st.strikethrough = _bool(attr(el, "strikethrough"), None)
    return st


def set_run_font(run, style: TextStyle) -> None:
    if style.font_size is not None:
        run.font.size = pt(style.font_size)
    if style.font_family:
        run.font.name = style.font_family
        # East Asian text uses a:ea; set it to the same family so Chinese
        # characters pick up the intended face instead of the theme default.
        rPr = run._r.get_or_add_rPr()
        for tag in ("a:latin", "a:ea", "a:cs"):
            e = rPr.find(qn(tag))
            if e is None:
                e = rPr.makeelement(qn(tag), {})
                rPr.append(e)
            e.set("typeface", style.font_family)
    if style.color:
        spec = parse_color(style.color)
        if spec and spec.rgb:
            run.font.color.rgb = spec.rgb
    if style.bold is not None:
        run.font.bold = style.bold
    if style.italic is not None:
        run.font.italic = style.italic
    if style.underline is not None:
        run.font.underline = style.underline
    if style.strikethrough is not None:
        run.font.strike = style.strikethrough


def set_paragraph_spacing(paragraph, line_spacing: Any, letter_spacing: Optional[float]) -> None:
    if line_spacing is not None:
        raw = str(line_spacing).strip()
        if raw.startswith("multiple:"):
            paragraph.line_spacing = float(raw.split(":", 1)[1])
        elif raw.startswith("fixed:"):
            paragraph.line_spacing = Pt(float(raw.split(":", 1)[1]))
        else:
            try:
                paragraph.line_spacing = float(raw)
            except ValueError:
                pass
    if letter_spacing is not None and letter_spacing != 0:
        # a:rPr @spc is in 1/100 pt; SML letterSpacing px == pt at 72dpi.
        pPr = paragraph._p.get_or_add_pPr()
        defRPr = pPr.find(qn("a:defRPr"))
        if defRPr is None:
            defRPr = pPr.makeelement(qn("a:defRPr"), {})
            pPr.append(defRPr)
        defRPr.set("spc", str(int(round(letter_spacing * 100))))


class RunBuilder:
    """Appends runs to a paragraph, flushing inline <br/> as <a:br/> runs."""

    def __init__(self, paragraph, style: TextStyle) -> None:
        self.paragraph = paragraph
        self.style = style
        self._pending_text: list[str] = []

    def flush(self) -> None:
        if not self._pending_text:
            return
        run = self.paragraph.add_run()
        run.text = "".join(self._pending_text)
        set_run_font(run, self.style)
        self._pending_text = []

    def text(self, value: str) -> None:
        self._pending_text.append(value)

    def br(self) -> None:
        self.flush()
        # <a:br/> is a direct child of <a:p>, not wrapped in a run
        self.paragraph._p.append(self.paragraph._p.makeelement(qn("a:br"), {}))

    def finish(self) -> None:
        self.flush()


def mixed_content(node) -> Any:
    """Yield an element's text + children + tails so no text node is lost."""
    if node.text:
        yield node.text
    for c in node:
        yield c
        if c.tail:
            yield c.tail


def render_inline_nodes(nodes, builder: RunBuilder, style: TextStyle, warnings: list[dict[str, Any]]) -> None:
    for node in nodes:
        if isinstance(node, str):
            builder.text(node)
            continue
        tag = local(node.tag)
        if tag == "br":
            builder.br()
        elif tag == "strong":
            st = TextStyle(); st.bold = True
            style.merge(st)
            render_inline_nodes(mixed_content(node), builder, style, warnings)
        elif tag == "em":
            st = TextStyle(); st.italic = True
            style.merge(st)
            render_inline_nodes(mixed_content(node), builder, style, warnings)
        elif tag == "u":
            st = TextStyle(); st.underline = True
            style.merge(st)
            render_inline_nodes(mixed_content(node), builder, style, warnings)
        elif tag == "del":
            st = TextStyle(); st.strikethrough = True
            style.merge(st)
            render_inline_nodes(mixed_content(node), builder, style, warnings)
        elif tag == "span":
            style.merge(span_style(node))
            render_inline_nodes(mixed_content(node), builder, style, warnings)
        elif tag == "a":
            render_inline_nodes(mixed_content(node), builder, style, warnings)
        elif tag == "formula":
            latex = "".join(node.itertext()).strip()
            builder.text(latex if latex else "(公式)")
            warnings.append(
                {"code": "formula_as_text", "message": "formula rendered as LaTeX source text; pre-render a local image when typesetting is required"}
            )
        elif tag in ("shadow", "outline", "field"):
            render_inline_nodes(mixed_content(node), builder, style, warnings)
        else:
            render_inline_nodes(mixed_content(node), builder, style, warnings)


def render_paragraph(paragraph, content_el: ET.Element, style: TextStyle, warnings: list[dict[str, Any]]) -> None:
    builder = RunBuilder(paragraph, style)
    render_inline_nodes(mixed_content(content_el), builder, style, warnings)
    builder.finish()


def fill_text_frame(
    text_frame,
    content_el: ET.Element,
    style: TextStyle,
    warnings: list[dict[str, Any]],
    default_align: str = "left",
) -> None:
    """Populate a text frame from an SML <content> element."""
    mode = attr(content_el, 'autoFit', 'no-auto-fit')
    text_frame.auto_size = getattr(MSO_AUTO_SIZE, AUTO_FIT_ENUM_NAMES[mode])
    text_frame.word_wrap = attr(content_el, "wrap", "true") != "false"
    va = attr(content_el, "verticalAlign")
    if va == "top":
        text_frame.vertical_anchor = MSO_ANCHOR.TOP
    elif va == "bottom":
        text_frame.vertical_anchor = MSO_ANCHOR.BOTTOM
    else:
        text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE

    align = attr(content_el, "textAlign") or default_align
    line_spacing = attr(content_el, "lineSpacing")
    letter_spacing = style.letter_spacing

    blocks: list[ET.Element] = []
    for c in content_el:
        tag = local(c.tag)
        if tag in ("p", "ul", "ol"):
            blocks.append(c)

    first = True
    for block in blocks:
        tag = local(block.tag)
        if tag == "p":
            if not first:
                paragraph = text_frame.add_paragraph()
            else:
                paragraph = text_frame.paragraphs[0]
                first = False
            paragraph.alignment = _pp_align(align)
            set_paragraph_spacing(paragraph, line_spacing, letter_spacing)
            render_paragraph(paragraph, block, style, warnings)
        else:  # ul / ol
            items = children(block, "li")
            for idx, li in enumerate(items):
                if not first:
                    paragraph = text_frame.add_paragraph()
                else:
                    paragraph = text_frame.paragraphs[0]
                    first = False
                paragraph.alignment = _pp_align(align)
                set_paragraph_spacing(paragraph, line_spacing, letter_spacing)
                bullet = "• " if tag == "ul" else f"{idx + 1}. "
                if attr(block, "start") is not None:
                    try:
                        bullet = f"{int(attr(block, 'start')) + idx}. "
                    except ValueError:
                        pass
                builder = RunBuilder(paragraph, style)
                builder.text(bullet)
                render_inline_nodes(mixed_content(li), builder, style, warnings)
                builder.finish()


def _pp_align(value: str) -> PP_ALIGN:
    return {
        "left": PP_ALIGN.LEFT,
        "center": PP_ALIGN.CENTER,
        "right": PP_ALIGN.RIGHT,
        "justify": PP_ALIGN.JUSTIFY,
        "dist": PP_ALIGN.JUSTIFY,
    }.get(value, PP_ALIGN.LEFT)


# --------------------------------------------------------------------------
# fills, borders
# --------------------------------------------------------------------------

def apply_shape_fill(shape, color: Optional[ColorSpec], shape_alpha: Optional[float], warnings=None) -> None:
    if color is None or color.kind == "transparent":
        shape.fill.background()
        return
    if color.kind == "linear-gradient" and color.gradient:
        try:
            fill = shape.fill
            fill.gradient()
            stops = fill.gradient_stops
            for idx, (rgb, pos) in enumerate(color.gradient):
                rgb_obj = _as_rgb(rgb)
                if idx < len(stops):
                    stops[idx].color.rgb = rgb_obj
                    stops[idx].position = pos
                else:
                    stops.insert(pos, rgb_obj)
            fill.gradient_angle = color.angle if color.angle is not None else 90.0
        except Exception as exc:
            # gradient API not available for this shape type; fall back to solid
            shape.fill.solid()
            shape.fill.fore_color.rgb = _as_rgb(color.gradient[0][0])
            if warnings is not None:
                warnings.append({'code': 'fill_approximated', 'message': f'Linear gradient fell back to first stop: {exc}'})
    elif color.kind == "radial-gradient" and color.gradient:
        # pptx has no radial gradient; approximate with the first stop
        shape.fill.solid()
        shape.fill.fore_color.rgb = _as_rgb(color.gradient[0][0])
    elif color.rgb is not None:
        shape.fill.solid()
        shape.fill.fore_color.rgb = color.rgb
    else:
        shape.fill.background()
        return
    alpha = alpha_of(shape_alpha) * alpha_of(color.alpha)
    if alpha < 1.0:
        _set_fill_alpha(shape, alpha)


def _as_rgb(rgb) -> RGBColor:
    if isinstance(rgb, RGBColor):
        return rgb
    return RGBColor(rgb[0], rgb[1], rgb[2])


def _set_fill_alpha(shape, alpha: float) -> None:
    """alpha 0..1 -> <a:alpha val="0..100000"/> on the shape's solid fill."""
    spPr = shape._element.spPr
    fill = spPr.find(qn("a:solidFill"))
    if fill is None:
        return
    srgb = fill.find(qn("a:srgbClr"))
    if srgb is None:
        return
    a = srgb.find(qn("a:alpha"))
    if a is None:
        a = srgb.makeelement(qn("a:alpha"), {})
        srgb.append(a)
    a.set("val", str(int(round(alpha * 100000))))


def apply_shape_border(shape, border_el: Optional[ET.Element]) -> None:
    if border_el is None:
        # SML defines an omitted border as no border. Native Office shapes,
        # however, inherit a theme outline unless we explicitly clear it.
        shape.line.fill.background()
        return
    color_spec = parse_color(attr(border_el, "color"))
    width = _num(attr(border_el, "width"), None)
    if color_spec is None and width is None:
        return
    line = shape.line
    if color_spec is None or color_spec.rgb is None:
        line.fill.background()
    else:
        line.color.rgb = color_spec.rgb
        alpha = alpha_of(color_spec.alpha)
        if alpha < 1.0:
            ln = shape._element.spPr.find(qn("a:ln"))
            if ln is not None:
                srgb = ln.find(qn("a:solidFill") + "/" + qn("a:srgbClr"))
                if srgb is None:
                    srgb = ln.find(qn("a:solidFill"))
                if srgb is not None:
                    srgb = srgb.find(qn("a:srgbClr"))
                if srgb is not None:
                    a = srgb.makeelement(qn("a:alpha"), {})
                    a.set("val", str(int(round(alpha * 100000))))
                    srgb.append(a)
    if width is not None:
        line.width = Pt(width)
    dash = LINE_DASH_MAP.get(str(attr(border_el, "dashArray", "solid")).lower())
    if dash:
        try:
            from pptx.enum.line import MSO_LINE
            line.dash_style = MSO_LINE.DASH if dash == "dash" else MSO_LINE.ROUND_DOT
        except Exception:
            pass


def add_arrowheads(shape, start_arrow: Optional[ET.Element], end_arrow: Optional[ET.Element]) -> None:
    ln = shape._element.spPr.find(qn("a:ln"))
    if ln is None:
        return
    for arrow_el, tail in ((start_arrow, "headEnd"), (end_arrow, "tailEnd")):
        if arrow_el is None:
            continue
        a_type = ARROW_TYPE_MAP.get(attr(arrow_el, "type", "none"))
        if not a_type or a_type == "none":
            continue
        end = ln.makeelement(qn(f"a:{tail}"), {})
        end.set("type", a_type)
        filled = attr(arrow_el, "type") not in ("empty-triangle", "empty-diamond", "empty-circle")
        end.set("w", "med")
        end.set("len", "med")
        ln.append(end)


# --------------------------------------------------------------------------
# background
# --------------------------------------------------------------------------

def apply_slide_background(slide, fill_el: Optional[ET.Element], warnings: list[dict[str, Any]]) -> None:
    if fill_el is None:
        return
    color_el = child(fill_el, "fillColor")
    spec = parse_color(attr(color_el, "color")) if color_el is not None else None
    if spec is None or spec.kind == "transparent":
        return
    cSld = slide._element.find(qn("p:cSld"))
    bg = cSld.find(qn("p:bg"))
    if bg is None:
        bg = cSld.makeelement(qn("p:bg"), {})
        cSld.insert(0, bg)
    bgPr = bg.find(qn("p:bgPr"))
    if bgPr is None:
        bgPr = bg.makeelement(qn("p:bgPr"), {})
        bg.append(bgPr)
    # remove any existing fill
    for tag in ("a:solidFill", "a:gradFill", "a:noFill"):
        e = bgPr.find(qn(tag))
        if e is not None:
            bgPr.remove(e)
    if spec.kind == "linear-gradient" and spec.gradient:
        grad = bgPr.makeelement(qn("a:gradFill"), {})
        gsLst = grad.makeelement(qn("a:gsLst"), {})
        for rgb, pos in spec.gradient:
            gs = gsLst.makeelement(qn("a:gs"), {"pos": str(int(round(pos * 100000)))})
            srgb = gs.makeelement(qn("a:srgbClr"), {"val": "%02X%02X%02X" % _rgb_tuple(rgb)})
            gs.append(srgb)
            gsLst.append(gs)
        grad.append(gsLst)
        ang = int(round((spec.angle if spec.angle is not None else 90.0) * 60000))
        lin = grad.makeelement(qn("a:lin"), {"ang": str(ang), "scaled": "1"})
        grad.append(lin)
        bgPr.append(grad)
    elif spec.rgb is not None or spec.kind == 'radial-gradient' and spec.gradient:
        solid = bgPr.makeelement(qn("a:solidFill"), {})
        rgb = spec.rgb if spec.rgb is not None else spec.gradient[0][0]
        srgb = solid.makeelement(qn("a:srgbClr"), {"val": "%02X%02X%02X" % _rgb_tuple(rgb)})
        solid.append(srgb)
        bgPr.append(solid)


def _rgb_tuple(rgb) -> tuple[int, int, int]:
    if isinstance(rgb, RGBColor):
        return rgb[0], rgb[1], rgb[2]
    return rgb[0], rgb[1], rgb[2]


# --------------------------------------------------------------------------
# element renderers
# --------------------------------------------------------------------------

def add_text_box(slide, el: ET.Element, style_defaults: dict[str, dict[str, Any]], warnings: list[dict[str, Any]]) -> None:
    x, y = px(attr(el, "topLeftX")), px(attr(el, "topLeftY"))
    w, h = px(attr(el, "width")), px(attr(el, "height"))
    box = slide.shapes.add_textbox(x, y, w, h)
    if attr(el, 'rotation') is not None:
        box.rotation = float(attr(el, 'rotation')) % 360.0
    content = child(el, "content")
    if content is None:
        return
    style = content_style(content, style_defaults)
    tf = box.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    fill_text_frame(tf, content, style, warnings, default_align="left")


def add_shape(slide, el: ET.Element, style_defaults: dict[str, dict[str, Any]], warnings: list[dict[str, Any]]) -> None:
    shape_type = attr(el, "type", "rect")
    x, y = px(attr(el, "topLeftX")), px(attr(el, "topLeftY"))
    w, h = px(attr(el, "width")), px(attr(el, "height"))
    mso = SHAPE_TYPE_MAP.get(shape_type)
    if mso is None:
        mso = MSO_SHAPE.RECTANGLE
        warnings.append({"code": "shape_type_approximated", "message": f"shape type '{shape_type}' approximated as rect"})
    shape = slide.shapes.add_shape(mso, x, y, w, h)
    shape.shadow.inherit = False
    fill_el = child(el, "fill")
    if fill_el is not None:
        color_el = child(fill_el, "fillColor")
        spec = parse_color(attr(color_el, "color")) if color_el is not None else None
        apply_shape_fill(shape, spec, _num(attr(el, "alpha"), None), warnings)
    else:
        shape.fill.background()
    apply_shape_border(shape, child(el, "border"))
    if attr(el, "presetHandlers") is not None and mso == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            radius = float(attr(el, "presetHandlers"))
            dim = min(float(attr(el, "width")), float(attr(el, "height")))
            clamped = max(0.0, min(radius, dim / 2.0))
            shape.adjustments[0] = clamped / dim if dim > 0 else 0.0
        except Exception:
            pass
    if attr(el, "rotation") is not None:
        shape.rotation = float(attr(el, "rotation")) % 360.0
    content = child(el, "content")
    if content is not None:
        style = content_style(content, style_defaults)
        tf = shape.text_frame
        tf.word_wrap = attr(content, "wrap", "true") != "false"
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        va = attr(content, "verticalAlign")
        if va == "top":
            tf.vertical_anchor = MSO_ANCHOR.TOP
        elif va == "bottom":
            tf.vertical_anchor = MSO_ANCHOR.BOTTOM
        else:
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        align = attr(content, "textAlign") or ("left" if shape_type == "text" else "center")
        blocks = [c for c in content if local(c.tag) in ("p", "ul", "ol")]
        first = True
        for block in blocks:
            if not first:
                paragraph = tf.add_paragraph()
            else:
                paragraph = tf.paragraphs[0]
                first = False
            paragraph.alignment = _pp_align(align)
            set_paragraph_spacing(paragraph, attr(content, "lineSpacing"), style.letter_spacing)
            if local(block.tag) == "p":
                render_paragraph(paragraph, block, style, warnings)
            else:
                items = children(block, "li")
                for idx, li in enumerate(items):
                    if not first:
                        paragraph = tf.add_paragraph()
                    else:
                        paragraph = tf.paragraphs[0]
                        first = False
                    paragraph.alignment = _pp_align(align)
                    set_paragraph_spacing(paragraph, attr(content, "lineSpacing"), style.letter_spacing)
                    bullet = "• " if local(block.tag) == "ul" else f"{idx + 1}. "
                    builder = RunBuilder(paragraph, style)
                    builder.text(bullet)
                    render_inline_nodes(mixed_content(li), builder, style, warnings)
                    builder.finish()


def add_line(slide, el: ET.Element, warnings: list[dict[str, Any]]) -> None:
    x1, y1 = px(attr(el, "startX")), px(attr(el, "startY"))
    x2, y2 = px(attr(el, "endX")), px(attr(el, "endY"))
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    conn.shadow.inherit = False
    border_el = child(el, "border")
    if border_el is not None:
        color_spec = parse_color(attr(border_el, "color"))
        width = _num(attr(border_el, "width"), None)
        if color_spec is None or color_spec.rgb is None:
            conn.line.fill.background()
        else:
            conn.line.color.rgb = color_spec.rgb
        if width is not None:
            conn.line.width = Pt(width)
    else:
        conn.line.color.rgb = RGBColor(0, 0, 0)
    add_arrowheads(conn, child(el, "startArrow"), child(el, "endArrow"))


def add_polyline(slide, el: ET.Element, warnings: list[dict[str, Any]]) -> None:
    ptype = attr(el, "type", "bent-connector2")
    x, y = px(attr(el, "topLeftX")), px(attr(el, "topLeftY"))
    w, h = px(attr(el, "width")), px(attr(el, "height"))
    conn_type = getattr(MSO_CONNECTOR, POLYLINE_ENUM_NAMES.get(ptype, 'CURVE' if 'curved' in ptype else 'ELBOW'))
    conn = slide.shapes.add_connector(conn_type, x, y, x + w, y + h)
    if ptype in POLYLINE_PRESETS:
        # The library factories otherwise use three-segment geometry for both.
        conn._element.find('.//' + qn('a:prstGeom')).set('prst', POLYLINE_PRESETS[ptype])
    conn.shadow.inherit = False
    border_el = child(el, "border")
    if border_el is not None:
        color_spec = parse_color(attr(border_el, "color"))
        width = _num(attr(border_el, "width"), None)
        if color_spec is None or color_spec.rgb is None:
            conn.line.fill.background()
        else:
            conn.line.color.rgb = color_spec.rgb
        if width is not None:
            conn.line.width = Pt(width)
    add_arrowheads(conn, child(el, "startArrow"), child(el, "endArrow"))
    if ptype not in POLYLINE_ENUM_NAMES:
        warnings.append({"code": "polyline_approximated", "message": f"{ptype} approximated as a single {conn_type.name.lower()} connector"})


def resolve_image_path(src: str, assets_dir: Optional[Path], media_manifest: Optional[dict[str, str]], warnings: list[dict[str, Any]]) -> Optional[Path]:
    if not src:
        return None
    if src.startswith("@") and len(src) > 1:
        src = src[1:]
    if src.startswith("http://") or src.startswith("https://"):
        warnings.append({"code": "remote_image_skipped", "message": f"remote image {src[:60]} is not fetched locally"})
        return None
    candidates: list[Path] = []
    if assets_dir:
        candidates.append(assets_dir / src)
    candidates.append(Path(src))
    for c in candidates:
        if c.exists() and c.is_file():
            return c
    # file_token from push-mode readback -> local path via manifest
    if media_manifest and src in media_manifest:
        m = Path(media_manifest[src])
        if m.exists():
            return m
    warnings.append({"code": "image_missing", "message": f"image not found locally: {src[:80]}"})
    return None


def _crop_factors(img_w: float, img_h: float, box_w: float, box_h: float, anchor: Optional[str]) -> tuple[float, float, float, float]:
    """Return (crop_left, crop_right, crop_top, crop_bottom) fractions."""
    if img_w <= 0 or img_h <= 0 or box_w <= 0 or box_h <= 0:
        return 0.0, 0.0, 0.0, 0.0
    src_ratio = img_w / img_h
    box_ratio = box_w / box_h
    if src_ratio > box_ratio:
        # image wider: crop horizontal, keep vertical fully
        keep = box_ratio / src_ratio
        excess = 1.0 - keep
        if anchor == "left":
            return 0.0, excess, 0.0, 0.0
        if anchor == "right":
            return excess, 0.0, 0.0, 0.0
        return excess / 2.0, excess / 2.0, 0.0, 0.0
    if src_ratio < box_ratio:
        keep = src_ratio / box_ratio
        excess = 1.0 - keep
        if anchor == "top":
            return 0.0, 0.0, 0.0, excess
        if anchor == "bottom":
            return 0.0, 0.0, excess, 0.0
        return 0.0, 0.0, excess / 2.0, excess / 2.0
    return 0.0, 0.0, 0.0, 0.0


def _partial_image_geometry(shape, width, height, radius, corners):
    """Translate model-selected corners to a native path, not a raster mask."""
    from pptx.oxml.xmlchemy import OxmlElement
    w, h, r = (int(round(v * 1000)) for v in (width, height, radius))
    selected = {'top': (True, True, False, False),
                'right': (False, True, True, False),
                'bottom': (False, False, True, True),
                'left': (True, False, False, True)}[corners]
    tl, tr, br, bl = (r if on else 0 for on in selected)
    geom = OxmlElement('a:custGeom')
    for tag in ('avLst', 'gdLst', 'ahLst', 'cxnLst'):
        geom.append(OxmlElement('a:' + tag))
    rect = OxmlElement('a:rect')
    for key, value in {'l':'0','t':'0','r':'r','b':'b'}.items():
        rect.set(key, value)
    geom.append(rect)
    paths = OxmlElement('a:pathLst'); geom.append(paths)
    path = OxmlElement('a:path'); path.set('w',str(w)); path.set('h',str(h)); paths.append(path)
    def command(tag, *points):
        node = OxmlElement('a:' + tag)
        for x, y in points:
            pt = OxmlElement('a:pt'); pt.set('x', str(round(x))); pt.set('y', str(round(y))); node.append(pt)
        path.append(node)
    k = 0.5522847498307936
    command('moveTo', (tl, 0)); command('lnTo', (w-tr, 0))
    if tr: command('cubicBezTo', (w-tr+k*tr,0), (w,tr-k*tr), (w,tr))
    command('lnTo', (w,h-br))
    if br: command('cubicBezTo', (w,h-br+k*br), (w-br+k*br,h), (w-br,h))
    command('lnTo', (bl,h))
    if bl: command('cubicBezTo', (bl-k*bl,h), (0,h-bl+k*bl), (0,h-bl))
    command('lnTo', (0,tl))
    if tl: command('cubicBezTo', (0,tl-k*tl), (tl-k*tl,0), (tl,0))
    command('close')
    sp = shape._element.spPr
    old = sp.find(qn('a:prstGeom')); sp.replace(old, geom)


def _native_image_shape(slide, path, x, y, w, h, mode, radius, corners, crop):
    from pptx.oxml.xmlchemy import OxmlElement
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL if mode == 'ellipse' else MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.shadow.inherit = False
    shape.fill.background()
    width, height = float(w) / px(1), float(h) / px(1)
    if mode != 'ellipse':
        if corners == 'all':
            shape.adjustments[0] = radius / min(width, height)
        else:
            _partial_image_geometry(shape, width, height, radius, corners)
    _, rid = slide.part.get_or_add_image_part(str(path))
    fill = OxmlElement('a:blipFill')
    blip = OxmlElement('a:blip'); blip.set(qn('r:embed'), rid); fill.append(blip)
    src = OxmlElement('a:srcRect')
    for key, value in zip(('l','r','t','b'), crop):
        src.set(key, str(round(value * 100000)))
    fill.append(src)
    stretch = OxmlElement('a:stretch'); stretch.append(OxmlElement('a:fillRect')); fill.append(stretch)
    sp = shape._element.spPr
    sp.replace(sp.find(qn('a:noFill')), fill)
    return shape


def add_image(slide, el: ET.Element, assets_dir: Optional[Path], media_manifest: Optional[dict[str, str]], warnings: list[dict[str, Any]]) -> None:
    from backend_capabilities import image_frame_spec
    crop_el = child(el, "crop")
    mode, radius, corners = image_frame_spec(crop_el, attr(el, 'width'), attr(el, 'height'))
    src = attr(el, "src")
    path = resolve_image_path(src or "", assets_dir, media_manifest, warnings)
    if path is None:
        return
    x, y = px(attr(el, "topLeftX")), px(attr(el, "topLeftY"))
    w, h = px(attr(el, "width")), px(attr(el, "height"))
    anchor = attr(crop_el, "anchor") if crop_el is not None else None
    with __import__("PIL").Image.open(path) as im:
        img_w, img_h = im.size
    cl, cr, ct, cb = _crop_factors(img_w, img_h, float(attr(el, "width")), float(attr(el, "height")), anchor)
    if mode == 'rect':
        pic = slide.shapes.add_picture(str(path), x, y, w, h)
        pic.crop_left, pic.crop_right, pic.crop_top, pic.crop_bottom = cl, cr, ct, cb
    else:
        pic = _native_image_shape(slide, path, x, y, w, h, mode, radius, corners, (cl, cr, ct, cb))
    border = child(el, 'border')
    if border is not None:
        border = ET.Element('border', {'color':'rgb(43,47,54)', 'width':'2', **border.attrib})
    apply_shape_border(pic, border)
    nv = pic._element.xpath('.//p:cNvPr')[0]
    if attr(el, 'id'):
        nv.set('name', attr(el, 'id'))
    nv.set('descr', attr(el, 'alt') or attr(el, 'desc') or '')
    if attr(el, "rotation") is not None:
        pic.rotation = float(attr(el, "rotation")) % 360.0


def add_icon(slide, el: ET.Element, assets_dir: Optional[Path], icon_renderer, warnings: list[dict[str, Any]]) -> None:
    icon_type = attr(el, "iconType", "")
    x, y = px(attr(el, "topLeftX")), px(attr(el, "topLeftY"))
    w, h = px(attr(el, "width")), px(attr(el, "height"))
    fill_el = child(el, "fill")
    color_spec = None
    if fill_el is not None:
        color_el = child(fill_el, "fillColor")
        color_spec = parse_color(attr(color_el, "color")) if color_el is not None else None
    color = color_spec.rgb if color_spec and color_spec.rgb else RGBColor(51, 51, 51)
    png = icon_renderer(icon_type, int(float(attr(el, "width"))), color, assets_dir) if icon_renderer else None
    if png and Path(png).exists():
        slide.shapes.add_picture(str(png), x, y, w, h)
    else:
        # offline placeholder: filled circle in the icon color
        shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, w, h)
        shape.fill.solid()
        shape.fill.fore_color.rgb = color
        shape.line.fill.background()
        shape.shadow.inherit = False
        warnings.append({"code": "icon_placeholder", "message": f"icon '{icon_type}' rendered as colored circle (no local PNG); pre-render or provide a local icon image"})


def add_table(slide, el: ET.Element, style_defaults: dict[str, dict[str, Any]], warnings: list[dict[str, Any]]) -> None:
    x, y = px(attr(el, "topLeftX")), px(attr(el, "topLeftY"))
    w, h = px(attr(el, "width")), px(attr(el, "height"))
    colgroup = child(el, "colgroup")
    col_els = children(colgroup, "col") if colgroup is not None else []
    col_count = max(len(col_els), 1)
    rows = children(el, "tr")
    row_count = max(len(rows), 1)
    table_shape = slide.shapes.add_table(row_count, col_count, x, y, w, h)
    table = table_shape.table
    table.first_row = False
    table.horz_banding = False
    # column widths: explicit col widths in px; remaining shared evenly
    explicit = [_num(c.get("width"), None) for c in col_els]
    total_explicit = sum(v for v in explicit if v is not None)
    remaining = max(0.0, float(attr(el, "width")) - total_explicit)
    free_cols = sum(1 for v in explicit if v is None) or 1
    for idx, col in enumerate(table.columns):
        v = explicit[idx] if idx < len(explicit) else None
        width_px = v if v is not None else (remaining / free_cols)
        col.width = px(width_px)
    # row heights
    explicit_rows = [_num(r.get("height"), None) for r in rows]
    total_explicit_rows = sum(v for v in explicit_rows if v is not None)
    remaining_rows = max(0.0, float(attr(el, "height")) - total_explicit_rows)
    free_rows = sum(1 for v in explicit_rows if v is None) or 1
    for idx, row in enumerate(table.rows):
        v = explicit_rows[idx] if idx < len(explicit_rows) else None
        height_px = v if v is not None else (remaining_rows / free_rows)
        row.height = px(height_px)
    for ri, tr in enumerate(rows):
        tds = children(tr, "td")
        for ci, td in enumerate(tds):
            if ri >= row_count or ci >= col_count:
                continue
            cell = table.cell(ri, ci)
            cell.margin_left = cell.margin_right = Pt(4)
            cell.margin_top = cell.margin_bottom = Pt(2)
            fill_el = child(td, "fill")
            if fill_el is not None:
                color_el = child(fill_el, "fillColor")
                spec = parse_color(attr(color_el, "color")) if color_el is not None else None
                if spec and spec.rgb:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = spec.rgb
                elif spec and spec.kind == "transparent":
                    cell.fill.background()
            else:
                cell.fill.background()
            content = child(td, "content")
            if content is not None:
                style = content_style(content, style_defaults)
                tf = cell.text_frame
                tf.word_wrap = attr(content, "wrap", "true") != "false"
                align = attr(content, "textAlign") or "center"
                blocks = [c for c in content if local(c.tag) in ("p", "ul", "ol")]
                first = True
                for block in blocks:
                    if not first:
                        paragraph = tf.add_paragraph()
                    else:
                        paragraph = tf.paragraphs[0]
                        first = False
                    paragraph.alignment = _pp_align(align)
                    set_paragraph_spacing(paragraph, attr(content, "lineSpacing"), style.letter_spacing)
                    if local(block.tag) == "p":
                        render_paragraph(paragraph, block, style, warnings)
                    else:
                        for idx, li in enumerate(children(block, "li")):
                            if not first:
                                paragraph = tf.add_paragraph()
                            else:
                                paragraph = tf.paragraphs[0]
                                first = False
                            builder = RunBuilder(paragraph, style)
                            builder.text("• ")
                            render_inline_nodes(mixed_content(li), builder, style, warnings)
                            builder.finish()


def _csv_values(value: str) -> list[str]:
    return [v.strip() for v in value.split(",") if v.strip() != ""]


def add_chart(slide, el: ET.Element, warnings: list[dict[str, Any]]) -> None:
    x, y = px(attr(el, "topLeftX")), px(attr(el, "topLeftY"))
    w, h = px(attr(el, "width")), px(attr(el, "height"))
    plot_area = child(el, "chartPlotArea")
    plot = child(plot_area, "chartPlot") if plot_area is not None else None
    chart_type_raw = attr(plot, "type", "column") if plot is not None else "column"
    data_el = child(el, "chartData")
    dims = dim_children(data_el) if data_el is not None else []
    if len(dims) < 2:
        warnings.append({"code": "chart_skipped", "message": "chart needs at least dim1 (categories) + dim2 (series)"})
        return
    categories = _csv_values("".join(children(dims[0], "chartField")[0].itertext())) if children(dims[0], "chartField") else []
    # donut detection: chartSectors innerRadius under chartPlot/chartSeriesList
    donut_inner: Optional[float] = None
    if plot is not None:
        series_list = child(plot, "chartSeriesList")
        if series_list is not None:
            for ser in children(series_list, "chartSeries"):
                sec = child(ser, "chartSectors")
                if sec is not None and attr(sec, "innerRadius") is not None:
                    try:
                        donut_inner = float(attr(sec, "innerRadius"))
                    except ValueError:
                        pass
                    break
    chart_type = CHART_TYPE_MAP.get(chart_type_raw)
    if chart_type is None:
        if chart_type_raw == "combo":
            chart_type = XL_CHART_TYPE.COLUMN_CLUSTERED
            warnings.append({"code": "chart_approximated", "message": "combo chart approximated as clustered column"})
        else:
            warnings.append({"code": "chart_approximated", "message": f"chart type '{chart_type_raw}' approximated as column"})
            chart_type = XL_CHART_TYPE.COLUMN_CLUSTERED
    if chart_type == XL_CHART_TYPE.PIE and donut_inner is not None:
        chart_type = XL_CHART_TYPE.DOUGHNUT
    cd = CategoryChartData()
    cd.categories = categories
    series_names: list[str] = []
    for dim in dims[1:]:
        fields = children(dim, "chartField")
        for field in fields:
            name = attr(field, "name") or f"系列{len(series_names) + 1}"
            values = _csv_values("".join(field.itertext()))
            nums: list[float] = []
            for v in values:
                try:
                    nums.append(float(v.replace(",", "").replace("%", "")))
                except ValueError:
                    warnings.append({"code": "chart_value_skipped", "message": f"non-numeric chart value '{v}'"})
            series_names.append(name)
            cd.add_series(name, nums)
    is_pie = chart_type in (XL_CHART_TYPE.PIE,)
    if is_pie:
        chart = slide.shapes.add_chart(XL_CHART_TYPE.PIE, x, y, w, h, cd).chart
    else:
        chart = slide.shapes.add_chart(chart_type, x, y, w, h, cd).chart
    # colors
    style_el = child(el, "chartStyle")
    theme_colors: list[RGBColor] = []
    if style_el is not None:
        color_theme = child(style_el, "chartColorTheme")
        if color_theme is not None:
            for c in children(color_theme, "color"):
                spec = parse_color(attr(c, "value"))
                if spec and spec.rgb:
                    theme_colors.append(spec.rgb)
    plots = chart.plots
    if plots:
        series = plots[0].series
        for idx, s in enumerate(series):
            if idx < len(theme_colors):
                try:
                    s.format.fill.solid()
                    s.format.fill.fore_color.rgb = theme_colors[idx]
                    s.format.line.fill.background()
                except Exception:
                    pass
    # labels
    labels = child(plot, "chartLabels") if plot is not None else None
    if labels is not None:
        try:
            plots[0].has_data_labels = True
            dls = plots[0].data_labels
            dls.show_category_name = _bool(attr(labels, "category"), False)
            dls.show_value = _bool(attr(labels, "value"), True)
            dls.show_percentage = _bool(attr(labels, "percentage"), False)
            if attr(labels, "fontSize") is not None:
                dls.font.size = pt(_num(attr(labels, "fontSize")))
            pos = attr(labels, "position")
            if pos == "outside" and is_pie:
                dls.position = XL_LABEL_POSITION.OUTSIDE_END
            elif pos == "inside":
                dls.position = XL_LABEL_POSITION.INSIDE_END
            elif pos in ("top", "auto"):
                pass  # python-pptx computes default placement
        except Exception:
            pass
    # legend
    legend_el = child(el, "chartLegend")
    if legend_el is not None:
        chart.has_legend = True
        pos = LEGEND_POSITION_MAP.get(attr(legend_el, "position", "right"))
        if pos:
            chart.legend.position = pos
        if attr(legend_el, "fontSize") is not None:
            chart.legend.font.size = pt(_num(attr(legend_el, "fontSize")))
    else:
        chart.has_legend = False
    # title
    title_el = child(el, "chartTitle")
    if title_el is not None:
        chart.has_title = True
        title_text = "".join(title_el.itertext()).strip()
        if title_text:
            chart.chart_title.text_frame.text = title_text
        if attr(title_el, "fontSize") is not None and chart.chart_title.text_frame.paragraphs:
            runs = chart.chart_title.text_frame.paragraphs[0].runs
            if runs:
                runs[0].font.size = pt(_num(attr(title_el, "fontSize")))
    else:
        chart.has_title = False
    # axis label font sizes + gridlines
    if plot_area is not None:
        axes = child(plot_area, "chartAxes")
        if axes is not None:
            for axis_el in children(axes, "chartAxis"):
                a_type = attr(axis_el, "type")
                label_el = child(axis_el, "chartLabel")
                if label_el is not None and attr(label_el, "fontSize") is not None:
                    size = pt(_num(attr(label_el, "fontSize")))
                    try:
                        if a_type == "x":
                            chart.category_axis.tick_labels.font.size = size
                        elif a_type == "y":
                            chart.value_axis.tick_labels.font.size = size
                    except Exception:
                        pass
                grid_el = child(axis_el, "chartGridLine")
                if grid_el is not None and attr(grid_el, "color") is not None:
                    spec = parse_color(attr(grid_el, "color"))
                    if spec and spec.rgb:
                        try:
                            chart.value_axis.major_gridlines.format.line.color.rgb = spec.rgb
                        except Exception:
                            pass
    # donut hole size
    if donut_inner is not None and plots:
        try:
            doughnut = plots[0]._element  # c:doughnutChart
            hole = doughnut.find(qn("c:holeSize"))
            if hole is None:
                hole = doughnut.makeelement(qn("c:holeSize"), {})
                doughnut.append(hole)
            hole.set("val", str(int(round(donut_inner * 100))))
        except Exception:
            pass


def add_note(slide, note_el: Optional[ET.Element], style_defaults: dict[str, dict[str, Any]], warnings: list[dict[str, Any]]) -> None:
    if note_el is None:
        return
    content = child(note_el, "content")
    if content is None:
        return
    notes = slide.notes_slide
    tf = notes.notes_text_frame
    style = content_style(content, style_defaults)
    blocks = [c for c in content if local(c.tag) in ("p", "ul", "ol")]
    first = True
    for block in blocks:
        if not first:
            paragraph = tf.add_paragraph()
        else:
            paragraph = tf.paragraphs[0]
            first = False
        if local(block.tag) == "p":
            render_paragraph(paragraph, block, style, warnings)
        else:
            for idx, li in enumerate(children(block, "li")):
                if not first:
                    paragraph = tf.add_paragraph()
                else:
                    paragraph = tf.paragraphs[0]
                    first = False
                builder = RunBuilder(paragraph, style)
                builder.text("• " if local(block.tag) == "ul" else f"{idx + 1}. ")
                render_inline_nodes(mixed_content(li), builder, style, warnings)
                builder.finish()


# --------------------------------------------------------------------------
# slide + presentation assembly
# --------------------------------------------------------------------------

def render_slide(
    prs: Presentation,
    slide_el: ET.Element,
    style_defaults: dict[str, dict[str, Any]],
    assets_dir: Optional[Path],
    media_manifest: Optional[dict[str, str]],
    icon_renderer,
    warnings: list[dict[str, Any]],
) -> dict[str, Any]:
    warnings.extend(capability_findings(slide_el))
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout
    style_el = child(slide_el, "style")
    if style_el is not None:
        fill_el = child(style_el, "fill")
        apply_slide_background(slide, fill_el, warnings)
    data_el = child(slide_el, "data")
    counts: dict[str, int] = {}
    if data_el is not None:
        for el in list(data_el):
            tag = local(el.tag)
            counts[tag] = counts.get(tag, 0) + 1
            if tag == "shape":
                if attr(el, "type", "rect") == "text":
                    add_text_box(slide, el, style_defaults, warnings)
                else:
                    add_shape(slide, el, style_defaults, warnings)
            elif tag == "line":
                add_line(slide, el, warnings)
            elif tag == "polyline":
                add_polyline(slide, el, warnings)
            elif tag == "img":
                add_image(slide, el, assets_dir, media_manifest, warnings)
            elif tag == "icon":
                add_icon(slide, el, assets_dir, icon_renderer, warnings)
            elif tag == "table":
                add_table(slide, el, style_defaults, warnings)
            elif tag == "chart":
                add_chart(slide, el, warnings)
            elif tag in ("undefined", "embed", "whiteboard"):
                warnings.append({"code": "element_skipped", "message": f"<{tag}> is server-only and was skipped"})
            else:
                warnings.append({"code": "element_skipped", "message": f"unsupported data element <{tag}> skipped"})
    add_note(slide, child(slide_el, "note"), style_defaults, warnings)
    return {"counts": counts}


def parse_theme(root: ET.Element) -> dict[str, dict[str, Any]]:
    defaults: dict[str, dict[str, Any]] = {}
    theme = child(root, "theme")
    if theme is None:
        return defaults
    text_styles = child(theme, "textStyles")
    if text_styles is None:
        return defaults
    for el in list(text_styles):
        tag = local(el.tag)
        if tag in TEXT_TYPE_DEFAULT_SIZE:
            entry: dict[str, Any] = {}
            if attr(el, "fontSize") is not None:
                entry["fontSize"] = _num(attr(el, "fontSize"))
            if attr(el, "fontFamily") is not None:
                entry["fontFamily"] = attr(el, "fontFamily")
            if attr(el, "fontColor") is not None:
                entry["fontColor"] = attr(el, "fontColor")
            defaults[tag] = entry
    return defaults


def render_presentation(
    prs: Presentation,
    root: ET.Element,
    assets_dir: Optional[Path],
    media_manifest: Optional[dict[str, str]],
    icon_renderer,
    warnings: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    width = _num(attr(root, "width"), CANVAS_W)
    height = _num(attr(root, "height"), CANVAS_H)
    prs.slide_width = px(width)
    prs.slide_height = px(height)
    style_defaults = parse_theme(root)
    report: list[dict[str, Any]] = []
    for slide_el in children(root, "slide"):
        report.append(
            render_slide(prs, slide_el, style_defaults, assets_dir, media_manifest, icon_renderer, warnings)
        )
    return report


def load_theme_from_dir(directory: Path) -> dict[str, dict[str, Any]]:
    theme_file = directory / "theme.xml"
    if theme_file.exists():
        try:
            root = ET.parse(str(theme_file)).getroot()
            return parse_theme(root)
        except ET.ParseError:
            pass
    return {}


def load_media_manifest(path: Optional[str]) -> Optional[dict[str, str]]:
    if not path:
        return None
    p = Path(path)
    if not p.exists():
        return None
    try:
        data = json.loads(p.read_text(encoding="utf-8"))
        if isinstance(data, dict):
            return {str(k): str(v) for k, v in data.items()}
    except (json.JSONDecodeError, OSError):
        pass
    return None


def collect_slide_files(directory: Path) -> list[Path]:
    def sort_key(p: Path) -> tuple[int, str]:
        m = re.search(r"(\d+)", p.name)
        return (int(m.group(1)) if m else 999999, p.name)
    return sorted([p for p in directory.glob("slide-*.xml") if p.is_file()], key=sort_key)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Convert Feishu SML XML to a local editable .pptx")
    parser.add_argument("--input", required=True, help="<presentation>.xml file or directory of slide-XX.xml files")
    parser.add_argument("--output", required=True, help="output .pptx path")
    parser.add_argument("--assets", default=None, help="directory holding local images referenced by <img src>")
    parser.add_argument("--media-manifest", default=None, help="JSON mapping file_token -> local image path")
    parser.add_argument("--title", default=None, help="deck title")
    parser.add_argument("--json", action="store_true", help="emit a JSON report on stdout")
    args = parser.parse_args(argv)

    warnings: list[dict[str, Any]] = []
    icon_renderer = None
    try:
        from iconpark_render import icon_renderer as _renderer  # type: ignore
        icon_renderer = _renderer
    except Exception:
        pass

    prs = Presentation()
    if args.title:
        prs.core_properties.title = args.title

    input_path = Path(args.input)
    if args.assets:
        assets_dir = Path(args.assets)
    elif input_path.is_dir() and (input_path / "assets").exists():
        assets_dir = input_path / "assets"
    elif input_path.is_dir():
        assets_dir = input_path
    else:
        assets_dir = input_path.parent
    media_manifest = load_media_manifest(args.media_manifest)
    report: list[dict[str, Any]] = []
    style_defaults: dict[str, dict[str, Any]] = {}

    if input_path.is_dir():
        slide_files = collect_slide_files(input_path)
        if not slide_files:
            print(json.dumps({"error": f"no slide-*.xml files in {input_path}"}))
            return 2
        width = CANVAS_W
        height = CANVAS_H
        for p in (input_path / "presentation.xml", input_path / "deck.xml"):
            if p.exists():
                try:
                    root = ET.parse(str(p)).getroot()
                    width = _num(attr(root, "width"), CANVAS_W)
                    height = _num(attr(root, "height"), CANVAS_H)
                    style_defaults = parse_theme(root)
                    break
                except ET.ParseError:
                    pass
        style_defaults.update(load_theme_from_dir(input_path))
        prs.slide_width = px(width)
        prs.slide_height = px(height)
        for sf in slide_files:
            try:
                root = ET.parse(str(sf)).getroot()
            except ET.ParseError as e:
                warnings.append({"code": "xml_parse_error", "message": f"{sf.name}: {e}"})
                continue
            if local(root.tag) == "presentation":
                report.extend(
                    render_presentation(prs, root, assets_dir, media_manifest, icon_renderer, warnings)
                )
            else:
                report.append(
                    render_slide(prs, root, style_defaults, assets_dir, media_manifest, icon_renderer, warnings)
                )
    else:
        if not input_path.exists():
            print(json.dumps({"error": f"input not found: {input_path}"}))
            return 2
        try:
            root = ET.parse(str(input_path)).getroot()
        except ET.ParseError as e:
            print(json.dumps({"error": f"XML parse error: {e}"}))
            return 2
        if local(root.tag) == "presentation":
            report = render_presentation(prs, root, assets_dir, media_manifest, icon_renderer, warnings)
        else:
            style_defaults = parse_theme(root)
            report.append(render_slide(prs, root, style_defaults, assets_dir, media_manifest, icon_renderer, warnings))

    out = Path(args.output)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))

    errors, advisory = partition(warnings)
    if args.json:
        print(
            json.dumps(
                {
                    "tool": "sml_to_pptx",
                    "output": str(out),
                    "slides": len(report),
                    "error_count": len(errors),
                    "errors": errors,
                    "warnings": advisory,
                    "slide_counts": report,
                },
                ensure_ascii=False,
                indent=2,
            )
        )
    return 1 if errors else 0


if __name__ == "__main__":
    sys.exit(main())
